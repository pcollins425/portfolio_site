"""Capacity dashboard — machine downtime and part cost for Ops / Finance management."""

from __future__ import annotations

import html
import os
import re
from calendar import monthrange
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

from app import mssql
from app import permission_catalog as cat
from app.auth_deps import require_area_read, require_demo_user
from app.casino_month_matrix import _ABBREV, _STATE_REGION

router = APIRouter(
    prefix="/api/capacity-dashboard",
    tags=["capacity-dashboard"],
    dependencies=[Depends(require_area_read(cat.CAPACITY_DASHBOARD_AREA))],
)

_MONTH_RE = re.compile(r"^\d{4}-\d{2}$")
_DAY_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_MAX_SPAN_DAYS = 365
_MAX_EXPORT_SPAN_DAYS = 3650  # 10 years — export-only ceiling beyond the dashboard cap

_INCLUDE = """
    (
      ISNULL(wo.hours, 0) > 0
      OR wo.downstart IS NOT NULL
      OR wo.downend IS NOT NULL
    )
"""

_BUCKET = "COALESCE(wo.downstart, wo.date_wo)"

ChartMode = Literal["trailing12", "range"]
ChartGrain = Literal["day", "week", "month"]


def _catalog() -> str:
    return (os.environ.get("MSSQL_DATABASE") or "dgs_application_db").strip()


def _query(sql: str, params=None) -> list[dict]:
    return mssql.query(sql, params=params, database=_catalog(), profile="field", load_env=False)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = html.unescape(str(value).strip())
    return text or None


def _day(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    return str(value)[:10]


def _num(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _parse_day(raw: str, name: str) -> date:
    text = (raw or "").strip()
    if not _DAY_RE.match(text):
        raise HTTPException(status_code=400, detail=f"{name} must be YYYY-MM-DD")
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{name} must be a valid date") from exc


def _parse_month(month: str | None) -> tuple[str, date, date]:
    today = date.today()
    raw = (month or "").strip() or today.strftime("%Y-%m")
    if not _MONTH_RE.match(raw):
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")
    year, mon = int(raw[:4]), int(raw[5:7])
    if mon < 1 or mon > 12:
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")
    start = date(year, mon, 1)
    end = date(year, mon, monthrange(year, mon)[1])
    return raw, start, end


def _is_full_month(start: date, end: date) -> bool:
    if start.year != end.year or start.month != end.month:
        return False
    last = monthrange(start.year, start.month)[1]
    return start.day == 1 and end.day == last


def _span_days(start: date, end: date) -> int:
    return (end - start).days + 1


def _month_key(day: date) -> str:
    return f"{day.year:04d}-{day.month:02d}"


def _add_months(day: date, months: int) -> date:
    year = day.year + (day.month - 1 + months) // 12
    month = (day.month - 1 + months) % 12 + 1
    return date(year, month, 1)


def _resolve_range(
    start_q: str | None,
    end_q: str | None,
    month_q: str | None,
    *,
    max_span_days: int | None = _MAX_SPAN_DAYS,
) -> tuple[date, date, str | None]:
    start_raw = (start_q or "").strip()
    end_raw = (end_q or "").strip()
    if start_raw or end_raw:
        if not start_raw or not end_raw:
            raise HTTPException(status_code=400, detail="start and end are both required")
        start = _parse_day(start_raw, "start")
        end = _parse_day(end_raw, "end")
        if end < start:
            raise HTTPException(status_code=400, detail="end must be on or after start")
        if max_span_days is not None and _span_days(start, end) > max_span_days:
            raise HTTPException(
                status_code=400,
                detail=f"range cannot exceed {max_span_days} days",
            )
        month_key = _month_key(start) if _is_full_month(start, end) else None
        return start, end, month_key
    month_key, start, end = _parse_month(month_q)
    return start, end, month_key


def _resolve_chart_mode(
    chart: str | None,
    *,
    is_full_month: bool,
) -> ChartMode:
    raw = (chart or "").strip().lower()
    if raw not in ("", "trailing12", "range"):
        raise HTTPException(status_code=400, detail="chart must be trailing12 or range")
    if not is_full_month:
        return "range"
    if raw == "range":
        return "range"
    return "trailing12"


def _adaptive_grain(start: date, end: date) -> ChartGrain:
    n = _span_days(start, end)
    if n <= 14:
        return "day"
    if n <= 84:
        return "week"
    return "month"


def _iso_monday(day: date) -> date:
    return day - timedelta(days=day.weekday())


def _empty_day() -> dict[str, float | int]:
    return {"downtime": 0.0, "wo_count": 0, "part_cost": 0.0}


def _daily_totals(
    start: date,
    end: date,
    *,
    filter_clause: str = "",
    filter_params: list[Any] | tuple[Any, ...] = (),
) -> dict[str, dict[str, float | int]]:
    out: dict[str, dict[str, float | int]] = {}
    params = (start, end, *filter_params)
    rows = _query(
        f"""
        SELECT
            CONVERT(char(10), {_BUCKET}, 23) AS day_key,
            ISNULL(SUM(CAST(wo.hours AS float)), 0) AS downtime,
            COUNT(*) AS wo_count
        {_base_joins(with_part_cost=False)}
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        GROUP BY CONVERT(char(10), {_BUCKET}, 23)
        """,
        params,
    )
    for row in rows:
        key = _text(row.get("day_key"))
        if not key:
            continue
        bucket = out.setdefault(key, _empty_day())
        bucket["downtime"] = round(_num(row.get("downtime")), 2)
        bucket["wo_count"] = int(row.get("wo_count") or 0)

    cost_rows = _query(
        f"""
        SELECT
            CONVERT(char(10), {_BUCKET}, 23) AS day_key,
            ISNULL(SUM(CAST(mat.ext_cost AS float)), 0) AS part_cost
        FROM projects.work_order_material AS mat
        INNER JOIN projects.work_orders AS wo ON wo.wo = mat.wo
        LEFT JOIN clients.casinos AS cas ON cas.reference_key = wo.casino_id
        LEFT JOIN clients.states AS st ON st.reference_key = cas.state_id
        LEFT JOIN clients.tribes AS tr ON tr.reference_key = cas.tribe_id
        LEFT JOIN inventory.assets AS ast ON ast.reference_key = wo.asset_id
        LEFT JOIN vendors.vendors AS ven ON ven.reference_key = ast.vendor_id
        LEFT JOIN vendors.cabinets AS cab ON cab.reference_key = ast.cabinet_id
        LEFT JOIN vendors.themes AS th ON th.reference_key = wo.theme_id
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        GROUP BY CONVERT(char(10), {_BUCKET}, 23)
        """,
        params,
    )
    for row in cost_rows:
        key = _text(row.get("day_key"))
        if not key:
            continue
        bucket = out.setdefault(key, _empty_day())
        bucket["part_cost"] = round(_num(row.get("part_cost")), 2)
    return out


def _sum_days(
    daily: dict[str, dict[str, float | int]],
    bucket_start: date,
    bucket_end: date,
) -> tuple[float, int, float]:
    downtime = 0.0
    wo_count = 0
    part_cost = 0.0
    cursor = bucket_start
    while cursor <= bucket_end:
        hit = daily.get(cursor.isoformat())
        if hit:
            downtime += float(hit["downtime"])
            wo_count += int(hit["wo_count"])
            part_cost += float(hit["part_cost"])
        cursor += timedelta(days=1)
    return round(downtime, 2), wo_count, round(part_cost, 2)


def _trend_point(
    *,
    key: str,
    label: str,
    bucket_start: date,
    bucket_end: date,
    downtime: float,
    wo_count: int,
    part_cost: float,
    incomplete: bool,
) -> dict[str, Any]:
    return {
        "key": key,
        "label": label,
        "start": bucket_start.isoformat(),
        "end": bucket_end.isoformat(),
        "downtime": downtime,
        "wo_count": wo_count,
        "part_cost": part_cost,
        "incomplete": incomplete,
        # Legacy field used by older UI builds.
        "month": key if len(key) == 7 else None,
    }


def _build_trailing12_trend(
    anchor_month_start: date,
    anchor_month_end: date,
    *,
    filter_clause: str = "",
    filter_params: list[Any] | tuple[Any, ...] = (),
) -> list[dict]:
    window_start = _add_months(anchor_month_start, -11)
    params = (window_start, anchor_month_end, *filter_params)
    rows = _query(
        f"""
        SELECT
            CONVERT(char(7), {_BUCKET}, 126) AS month_key,
            ISNULL(SUM(CAST(wo.hours AS float)), 0) AS downtime,
            COUNT(*) AS wo_count
        {_base_joins(with_part_cost=False)}
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        GROUP BY CONVERT(char(7), {_BUCKET}, 126)
        ORDER BY month_key
        """,
        params,
    )
    by_month: dict[str, dict[str, float | int]] = {}
    for row in rows:
        key = _text(row.get("month_key"))
        if not key:
            continue
        by_month[key] = {
            "downtime": round(_num(row.get("downtime")), 2),
            "wo_count": int(row.get("wo_count") or 0),
            "part_cost": 0.0,
        }

    cost_rows = _query(
        f"""
        SELECT
            CONVERT(char(7), {_BUCKET}, 126) AS month_key,
            ISNULL(SUM(CAST(mat.ext_cost AS float)), 0) AS part_cost
        FROM projects.work_order_material AS mat
        INNER JOIN projects.work_orders AS wo ON wo.wo = mat.wo
        LEFT JOIN clients.casinos AS cas ON cas.reference_key = wo.casino_id
        LEFT JOIN clients.states AS st ON st.reference_key = cas.state_id
        LEFT JOIN clients.tribes AS tr ON tr.reference_key = cas.tribe_id
        LEFT JOIN inventory.assets AS ast ON ast.reference_key = wo.asset_id
        LEFT JOIN vendors.vendors AS ven ON ven.reference_key = ast.vendor_id
        LEFT JOIN vendors.cabinets AS cab ON cab.reference_key = ast.cabinet_id
        LEFT JOIN vendors.themes AS th ON th.reference_key = wo.theme_id
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        GROUP BY CONVERT(char(7), {_BUCKET}, 126)
        """,
        params,
    )
    for row in cost_rows:
        key = _text(row.get("month_key"))
        if not key:
            continue
        bucket = by_month.setdefault(key, _empty_day())
        bucket["part_cost"] = round(_num(row.get("part_cost")), 2)

    out: list[dict] = []
    year, mon = window_start.year, window_start.month
    for _ in range(12):
        key = f"{year:04d}-{mon:02d}"
        last = monthrange(year, mon)[1]
        b_start = date(year, mon, 1)
        b_end = date(year, mon, last)
        hit = by_month.get(key, _empty_day())
        out.append(
            _trend_point(
                key=key,
                label=f"{mon:02d}",
                bucket_start=b_start,
                bucket_end=b_end,
                downtime=float(hit["downtime"]),
                wo_count=int(hit["wo_count"]),
                part_cost=float(hit["part_cost"]),
                incomplete=False,
            )
        )
        mon += 1
        if mon > 12:
            mon = 1
            year += 1
    return out


def _build_range_trend(
    start: date,
    end: date,
    *,
    filter_clause: str = "",
    filter_params: list[Any] | tuple[Any, ...] = (),
) -> tuple[ChartGrain, list[dict]]:
    grain = _adaptive_grain(start, end)
    daily = _daily_totals(
        start, end, filter_clause=filter_clause, filter_params=filter_params
    )
    points: list[dict] = []

    if grain == "day":
        cursor = start
        while cursor <= end:
            hit = daily.get(cursor.isoformat(), _empty_day())
            points.append(
                _trend_point(
                    key=cursor.isoformat(),
                    label=f"{cursor.month}/{cursor.day}",
                    bucket_start=cursor,
                    bucket_end=cursor,
                    downtime=float(hit["downtime"]),
                    wo_count=int(hit["wo_count"]),
                    part_cost=float(hit["part_cost"]),
                    incomplete=False,
                )
            )
            cursor += timedelta(days=1)
        return grain, points

    if grain == "week":
        week_start = _iso_monday(start)
        while week_start <= end:
            week_end = week_start + timedelta(days=6)
            clipped_start = max(week_start, start)
            clipped_end = min(week_end, end)
            if clipped_start <= clipped_end:
                downtime, wo_count, part_cost = _sum_days(daily, clipped_start, clipped_end)
                incomplete = clipped_start != week_start or clipped_end != week_end
                label = (
                    f"{clipped_start.month}/{clipped_start.day}"
                    if clipped_start == clipped_end
                    else f"{clipped_start.month}/{clipped_start.day}–{clipped_end.month}/{clipped_end.day}"
                )
                points.append(
                    _trend_point(
                        key=week_start.isoformat(),
                        label=label,
                        bucket_start=clipped_start,
                        bucket_end=clipped_end,
                        downtime=downtime,
                        wo_count=wo_count,
                        part_cost=part_cost,
                        incomplete=incomplete,
                    )
                )
            week_start += timedelta(days=7)
        return grain, points

    # month
    year, mon = start.year, start.month
    while date(year, mon, 1) <= end:
        last = monthrange(year, mon)[1]
        month_start = date(year, mon, 1)
        month_end = date(year, mon, last)
        clipped_start = max(month_start, start)
        clipped_end = min(month_end, end)
        downtime, wo_count, part_cost = _sum_days(daily, clipped_start, clipped_end)
        incomplete = clipped_start != month_start or clipped_end != month_end
        key = f"{year:04d}-{mon:02d}"
        points.append(
            _trend_point(
                key=key,
                label=f"{mon:02d}",
                bucket_start=clipped_start,
                bucket_end=clipped_end,
                downtime=downtime,
                wo_count=wo_count,
                part_cost=part_cost,
                incomplete=incomplete,
            )
        )
        mon += 1
        if mon > 12:
            mon = 1
            year += 1
    return grain, points


def _region_case_sql(expr: str = "st.state") -> str:
    whens = [f"WHEN N'{name}' THEN N'{region}'" for name, region in _STATE_REGION.items()]
    for abbr, full in _ABBREV.items():
        whens.append(f"WHEN N'{abbr}' THEN N'{_STATE_REGION[full.lower()]}'")
    return f"CASE LOWER(RTRIM({expr})) " + " ".join(whens) + " ELSE N'Unassigned' END"


def _ranked(rows: list[dict], *, limit: int | None = None) -> list[dict]:
    out = []
    selected = rows if limit is None else rows[:limit]
    for row in selected:
        key = _text(row.get("dim_key")) or "Unassigned"
        label = _text(row.get("dim_label")) or key
        out.append(
            {
                "key": key,
                "label": label,
                "downtime": round(_num(row.get("downtime")), 2),
                "wo_count": int(row.get("wo_count") or 0),
                "asset_count": int(row.get("asset_count") or 0),
                "part_cost": round(_num(row.get("part_cost")), 2),
            }
        )
    return out


def _base_joins(*, with_part_cost: bool = True) -> str:
    part = ""
    if with_part_cost:
        part = """
        LEFT JOIN (
            SELECT wo AS wo_key, SUM(CAST(ext_cost AS float)) AS part_cost
            FROM projects.work_order_material
            GROUP BY wo
        ) AS wpc ON wpc.wo_key = wo.wo
        """
    return f"""
        FROM projects.work_orders AS wo
        LEFT JOIN clients.casinos AS cas ON cas.reference_key = wo.casino_id
        LEFT JOIN clients.states AS st ON st.reference_key = cas.state_id
        LEFT JOIN clients.tribes AS tr ON tr.reference_key = cas.tribe_id
        LEFT JOIN inventory.assets AS ast ON ast.reference_key = wo.asset_id
        LEFT JOIN vendors.vendors AS ven ON ven.reference_key = ast.vendor_id
        LEFT JOIN vendors.cabinets AS cab ON cab.reference_key = ast.cabinet_id
        LEFT JOIN vendors.themes AS th ON th.reference_key = wo.theme_id
        {part}
    """


_FILTER_DIMS = ("region", "tribe", "state", "casino", "vendor", "cabinet", "theme")


def _parse_csv_param(raw: str | None) -> list[str]:
    if not raw:
        return []
    seen: set[str] = set()
    out: list[str] = []
    for part in str(raw).split(","):
        value = part.strip()
        if not value or value in seen:
            continue
        seen.add(value)
        out.append(value)
    return out


def _collect_filters(
    *,
    region: str | None = None,
    tribe: str | None = None,
    state: str | None = None,
    casino: str | None = None,
    vendor: str | None = None,
    cabinet: str | None = None,
    theme: str | None = None,
) -> dict[str, list[str]]:
    raw = {
        "region": region,
        "tribe": tribe,
        "state": state,
        "casino": casino,
        "vendor": vendor,
        "cabinet": cabinet,
        "theme": theme,
    }
    return {key: _parse_csv_param(value) for key, value in raw.items() if _parse_csv_param(value)}


def _dim_expr(dim: str, region_sql: str) -> str:
    mapping = {
        "region": region_sql,
        "tribe": "COALESCE(NULLIF(LTRIM(RTRIM(cas.tribe_id)), N''), N'Unassigned')",
        "state": "COALESCE(NULLIF(LTRIM(RTRIM(st.reference_key)), N''), N'Unassigned')",
        "casino": "COALESCE(NULLIF(LTRIM(RTRIM(wo.casino_id)), N''), N'Unassigned')",
        "vendor": "COALESCE(NULLIF(LTRIM(RTRIM(ast.vendor_id)), N''), N'Unassigned')",
        "cabinet": "COALESCE(NULLIF(LTRIM(RTRIM(ast.cabinet_id)), N''), N'Unassigned')",
        "theme": "COALESCE(NULLIF(LTRIM(RTRIM(wo.theme_id)), N''), N'Unassigned')",
    }
    if dim not in mapping:
        raise HTTPException(status_code=400, detail=f"unknown filter dim: {dim}")
    return mapping[dim]


def _dim_label_expr(dim: str, region_sql: str) -> str:
    mapping = {
        "region": region_sql,
        "tribe": "COALESCE(NULLIF(LTRIM(RTRIM(tr.tribe_name)), N''), N'Unassigned')",
        "state": "COALESCE(NULLIF(LTRIM(RTRIM(st.state)), N''), N'Unassigned')",
        "casino": "COALESCE(NULLIF(LTRIM(RTRIM(cas.casino_name)), N''), N'Unassigned')",
        "vendor": "COALESCE(NULLIF(LTRIM(RTRIM(ven.vendor_name)), N''), N'Unassigned')",
        "cabinet": "COALESCE(NULLIF(LTRIM(RTRIM(cab.cabinet_name)), N''), N'Unassigned')",
        "theme": (
            "COALESCE(NULLIF(LTRIM(RTRIM(th.theme_name)), N''), "
            "NULLIF(LTRIM(RTRIM(wo.comp_desc)), N''), N'Unassigned')"
        ),
    }
    return mapping[dim]


def _filter_sql(filters: dict[str, list[str]], region_sql: str) -> tuple[str, list[Any]]:
    if not filters:
        return "", []
    parts: list[str] = []
    params: list[Any] = []
    for dim in _FILTER_DIMS:
        values = filters.get(dim) or []
        if not values:
            continue
        expr = _dim_expr(dim, region_sql)
        placeholders = ", ".join(["%s"] * len(values))
        parts.append(f"{expr} IN ({placeholders})")
        params.extend(values)
    if not parts:
        return "", []
    return " AND " + " AND ".join(parts), params


def _query_filter_params(
    region: str | None = None,
    tribe: str | None = None,
    state: str | None = None,
    casino: str | None = None,
    vendor: str | None = None,
    cabinet: str | None = None,
    theme: str | None = None,
) -> tuple[dict[str, list[str]], str, list[Any]]:
    filters = _collect_filters(
        region=region,
        tribe=tribe,
        state=state,
        casino=casino,
        vendor=vendor,
        cabinet=cabinet,
        theme=theme,
    )
    clause, params = _filter_sql(filters, _region_case_sql())
    return filters, clause, params


@router.get("/filter-options")
def capacity_filter_options(
    start: Annotated[str | None, Query()] = None,
    end: Annotated[str | None, Query()] = None,
    month: Annotated[str | None, Query()] = None,
    user: Annotated[dict[str, Any] | None, Depends(require_demo_user)] = None,
):
    del user
    range_start, range_end, _month_key = _resolve_range(start, end, month)
    region_sql = _region_case_sql()
    options: dict[str, list[dict[str, str]]] = {}
    for dim in _FILTER_DIMS:
        key_expr = _dim_expr(dim, region_sql)
        label_expr = _dim_label_expr(dim, region_sql)
        rows = _query(
            f"""
            SELECT DISTINCT
                {key_expr} AS value_key,
                {label_expr} AS value_label
            {_base_joins(with_part_cost=False)}
            WHERE {_BUCKET} >= %s
              AND {_BUCKET} < DATEADD(day, 1, %s)
            ORDER BY value_label
            """,
            (range_start, range_end),
        )
        seen: set[str] = set()
        items: list[dict[str, str]] = []
        for row in rows:
            value = _text(row.get("value_key")) or "Unassigned"
            if value in seen:
                continue
            seen.add(value)
            items.append(
                {
                    "value": value,
                    "label": _text(row.get("value_label")) or value,
                }
            )
        options[dim] = items
    return {
        "start": range_start.isoformat(),
        "end": range_end.isoformat(),
        "options": options,
    }


@router.get("")
def capacity_dashboard(
    start: Annotated[str | None, Query()] = None,
    end: Annotated[str | None, Query()] = None,
    month: Annotated[str | None, Query()] = None,
    chart: Annotated[str | None, Query()] = None,
    region: Annotated[str | None, Query()] = None,
    tribe: Annotated[str | None, Query()] = None,
    state: Annotated[str | None, Query()] = None,
    casino: Annotated[str | None, Query()] = None,
    vendor: Annotated[str | None, Query()] = None,
    cabinet: Annotated[str | None, Query()] = None,
    theme: Annotated[str | None, Query()] = None,
    user: Annotated[dict[str, Any] | None, Depends(require_demo_user)] = None,
):
    del user
    range_start, range_end, month_key = _resolve_range(start, end, month)
    is_full = _is_full_month(range_start, range_end)
    chart_mode = _resolve_chart_mode(chart, is_full_month=is_full)
    region_sql = _region_case_sql()
    active_filters, filter_clause, filter_params = _query_filter_params(
        region=region,
        tribe=tribe,
        state=state,
        casino=casino,
        vendor=vendor,
        cabinet=cabinet,
        theme=theme,
    )
    range_params = (range_start, range_end, *filter_params)

    kpi_rows = _query(
        f"""
        SELECT
            ISNULL(SUM(CAST(wo.hours AS float)), 0) AS downtime,
            COUNT(*) AS wo_count,
            COUNT(DISTINCT COALESCE(NULLIF(LTRIM(RTRIM(wo.asset_id)), N''), wo.compid)) AS asset_count,
            AVG(CAST(NULLIF(wo.hours, 0) AS float)) AS avg_downtime
        {_base_joins()}
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        """,
        range_params,
    )
    part_kpi = _query(
        f"""
        SELECT ISNULL(SUM(CAST(mat.ext_cost AS float)), 0) AS part_cost
        FROM projects.work_order_material AS mat
        INNER JOIN projects.work_orders AS wo ON wo.wo = mat.wo
        LEFT JOIN clients.casinos AS cas ON cas.reference_key = wo.casino_id
        LEFT JOIN clients.states AS st ON st.reference_key = cas.state_id
        LEFT JOIN clients.tribes AS tr ON tr.reference_key = cas.tribe_id
        LEFT JOIN inventory.assets AS ast ON ast.reference_key = wo.asset_id
        LEFT JOIN vendors.vendors AS ven ON ven.reference_key = ast.vendor_id
        LEFT JOIN vendors.cabinets AS cab ON cab.reference_key = ast.cabinet_id
        LEFT JOIN vendors.themes AS th ON th.reference_key = wo.theme_id
        WHERE {_INCLUDE}
          AND {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        """,
        range_params,
    )
    kpi = kpi_rows[0] if kpi_rows else {}
    downtime = _num(kpi.get("downtime"))
    wo_count = int(kpi.get("wo_count") or 0)
    avg = _num(kpi.get("avg_downtime")) if wo_count else 0.0

    if chart_mode == "trailing12":
        chart_grain: ChartGrain = "month"
        trend = _build_trailing12_trend(
            range_start,
            range_end,
            filter_clause=filter_clause,
            filter_params=filter_params,
        )
    else:
        chart_grain, trend = _build_range_trend(
            range_start,
            range_end,
            filter_clause=filter_clause,
            filter_params=filter_params,
        )

    def breakout(dim_key_sql: str, dim_label_sql: str) -> list[dict]:
        # All work orders in the selected period (not downtime-only), so part cost
        # and zero-downtime WOs still appear. Default order is downtime desc.
        rows = _query(
            f"""
            SELECT
                {dim_key_sql} AS dim_key,
                {dim_label_sql} AS dim_label,
                ISNULL(SUM(CAST(wo.hours AS float)), 0) AS downtime,
                COUNT(*) AS wo_count,
                COUNT(DISTINCT COALESCE(NULLIF(LTRIM(RTRIM(wo.asset_id)), N''), wo.compid)) AS asset_count,
                ISNULL(SUM(ISNULL(wpc.part_cost, 0)), 0) AS part_cost
            {_base_joins()}
            WHERE {_BUCKET} >= %s
              AND {_BUCKET} < DATEADD(day, 1, %s)
              {filter_clause}
            GROUP BY {dim_key_sql}, {dim_label_sql}
            ORDER BY downtime DESC, wo_count DESC
            """,
            range_params,
        )
        return _ranked(rows)

    geo = {
        "region": breakout(region_sql, region_sql),
        "state": breakout(
            "COALESCE(NULLIF(LTRIM(RTRIM(st.reference_key)), N''), N'Unassigned')",
            "COALESCE(NULLIF(LTRIM(RTRIM(st.state)), N''), N'Unassigned')",
        ),
        "casino": breakout(
            "COALESCE(NULLIF(LTRIM(RTRIM(wo.casino_id)), N''), N'Unassigned')",
            "COALESCE(NULLIF(LTRIM(RTRIM(cas.casino_name)), N''), N'Unassigned')",
        ),
    }
    product = {
        "vendor": breakout(
            "COALESCE(NULLIF(LTRIM(RTRIM(ast.vendor_id)), N''), N'Unassigned')",
            "COALESCE(NULLIF(LTRIM(RTRIM(ven.vendor_name)), N''), N'Unassigned')",
        ),
        "cabinet": breakout(
            "COALESCE(NULLIF(LTRIM(RTRIM(ast.cabinet_id)), N''), N'Unassigned')",
            "COALESCE(NULLIF(LTRIM(RTRIM(cab.cabinet_name)), N''), N'Unassigned')",
        ),
        "theme": breakout(
            "COALESCE(NULLIF(LTRIM(RTRIM(wo.theme_id)), N''), N'Unassigned')",
            "COALESCE(NULLIF(LTRIM(RTRIM(th.theme_name)), N''), NULLIF(LTRIM(RTRIM(wo.comp_desc)), N''), N'Unassigned')",
        ),
    }

    grain_label = {"day": "by day", "week": "by week", "month": "by month"}[chart_grain]
    if chart_mode == "trailing12":
        chart_title = "Part cost & downtime by month"
        chart_subtitle = "trailing 12 mo"
    else:
        chart_title = f"Part cost & downtime {grain_label}"
        chart_subtitle = grain_label

    return {
        "month": month_key,
        "start": range_start.isoformat(),
        "end": range_end.isoformat(),
        "month_start": range_start.isoformat(),
        "month_end": range_end.isoformat(),
        "is_full_month": is_full,
        "chart_mode": chart_mode,
        "chart_grain": chart_grain,
        "chart_title": chart_title,
        "chart_subtitle": chart_subtitle,
        "max_span_days": _MAX_SPAN_DAYS,
        "filters": active_filters,
        "kpis": {
            "downtime": round(downtime, 2),
            "wo_count": wo_count,
            "asset_count": int(kpi.get("asset_count") or 0),
            "avg_downtime": round(avg, 2),
            "part_cost": round(_num(part_kpi[0].get("part_cost") if part_kpi else 0), 2),
        },
        "trend": trend,
        "geo": geo,
        "product": product,
    }


_DRILL_DIMS = frozenset({"region", "state", "casino", "vendor", "cabinet", "theme"})


@router.get("/drill")
def capacity_drill(
    start: Annotated[str | None, Query()] = None,
    end: Annotated[str | None, Query()] = None,
    month: Annotated[str | None, Query()] = None,
    dim: Annotated[str, Query()] = "",
    key: Annotated[str, Query()] = "",
    region: Annotated[str | None, Query()] = None,
    tribe: Annotated[str | None, Query()] = None,
    state: Annotated[str | None, Query()] = None,
    casino: Annotated[str | None, Query()] = None,
    vendor: Annotated[str | None, Query()] = None,
    cabinet: Annotated[str | None, Query()] = None,
    theme: Annotated[str | None, Query()] = None,
    user: Annotated[dict[str, Any] | None, Depends(require_demo_user)] = None,
):
    del user
    range_start, range_end, month_key = _resolve_range(start, end, month)
    dim = (dim or "").strip().lower()
    key = (key or "").strip() or "Unassigned"
    # Blank / "range" = all work orders in the date window (chart title / bucket click).
    if dim and dim not in _DRILL_DIMS and dim != "range":
        raise HTTPException(
            status_code=400,
            detail="dim must be region, state, casino, vendor, cabinet, theme, range, or omitted",
        )

    region_sql = _region_case_sql()
    drill_exprs = {
        "region": region_sql,
        "state": "COALESCE(NULLIF(LTRIM(RTRIM(st.reference_key)), N''), N'Unassigned')",
        "casino": "COALESCE(NULLIF(LTRIM(RTRIM(wo.casino_id)), N''), N'Unassigned')",
        "vendor": "COALESCE(NULLIF(LTRIM(RTRIM(ast.vendor_id)), N''), N'Unassigned')",
        "cabinet": "COALESCE(NULLIF(LTRIM(RTRIM(ast.cabinet_id)), N''), N'Unassigned')",
        "theme": "COALESCE(NULLIF(LTRIM(RTRIM(wo.theme_id)), N''), N'Unassigned')",
    }
    _active_filters, filter_clause, filter_params = _query_filter_params(
        region=region,
        tribe=tribe,
        state=state,
        casino=casino,
        vendor=vendor,
        cabinet=cabinet,
        theme=theme,
    )
    dim_filter = ""
    params: list[Any] = [range_start, range_end, *filter_params]
    if dim in _DRILL_DIMS:
        dim_filter = f" AND {drill_exprs[dim]} = %s"
        params.append(key)

    rows = _query(
        f"""
        SELECT TOP 300
            wo.wo,
            wo.date_wo,
            wo.hours,
            wo.downstart,
            wo.downend,
            wo.property,
            wo.casino_id,
            wo.asset_id,
            wo.theme_id,
            ast.vendor_id,
            ast.cabinet_id,
            COALESCE(ast.serial_number, wo.compid) AS serial,
            ven.vendor_name,
            cab.cabinet_name,
            COALESCE(th.theme_name, wo.comp_desc) AS theme_name,
            cas.casino_name,
            ISNULL(wpc.part_cost, 0) AS part_cost
        {_base_joins()}
        WHERE {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
          {dim_filter}
        ORDER BY ISNULL(wo.hours, 0) DESC, wo.downstart DESC, wo.wo DESC
        """,
        tuple(params),
    )
    return {
        "month": month_key,
        "start": range_start.isoformat(),
        "end": range_end.isoformat(),
        "dim": dim or "range",
        "key": key if dim in _DRILL_DIMS else "range",
        "orders": [
            {
                "wo": _text(row.get("wo")),
                "date_wo": _day(row.get("date_wo")),
                "hours": round(_num(row.get("hours")), 2) if row.get("hours") is not None else None,
                "downstart": _day(row.get("downstart")),
                "downend": _day(row.get("downend")),
                "property": _text(row.get("casino_name") or "Unassigned"),
                "casino_id": _text(row.get("casino_id")),
                "asset_id": _text(row.get("asset_id")),
                "theme_id": _text(row.get("theme_id")),
                "vendor_id": _text(row.get("vendor_id")),
                "cabinet_id": _text(row.get("cabinet_id")),
                "serial": _text(row.get("serial")),
                "vendor": _text(row.get("vendor_name")),
                "cabinet": _text(row.get("cabinet_name")),
                "theme": _text(row.get("theme_name")),
                "part_cost": round(_num(row.get("part_cost")), 2),
            }
            for row in rows
        ],
    }


_EXPORT_HEADERS = (
    "WO #",
    "WO date",
    "Downtime (days)",
    "Part cost",
    "Down start",
    "Down end",
    "Region",
    "Tribe",
    "State",
    "Casino",
    "Serial",
    "Vendor",
    "Cabinet",
    "Theme",
)


def _line_rows_for_export(
    range_start: date,
    range_end: date,
    *,
    filter_clause: str,
    filter_params: list[Any] | tuple[Any, ...],
) -> list[dict[str, Any]]:
    region_sql = _region_case_sql()
    rows = _query(
        f"""
        SELECT
            wo.wo,
            wo.date_wo,
            wo.hours,
            wo.downstart,
            wo.downend,
            {region_sql} AS region_name,
            COALESCE(NULLIF(LTRIM(RTRIM(tr.tribe_name)), N''), N'Unassigned') AS tribe_name,
            COALESCE(NULLIF(LTRIM(RTRIM(st.state)), N''), N'Unassigned') AS state_name,
            COALESCE(NULLIF(LTRIM(RTRIM(cas.casino_name)), N''), N'Unassigned') AS casino_name,
            COALESCE(ast.serial_number, wo.compid) AS serial,
            ven.vendor_name,
            cab.cabinet_name,
            COALESCE(th.theme_name, wo.comp_desc) AS theme_name,
            ISNULL(wpc.part_cost, 0) AS part_cost
        {_base_joins()}
        WHERE {_BUCKET} >= %s
          AND {_BUCKET} < DATEADD(day, 1, %s)
          {filter_clause}
        ORDER BY ISNULL(wo.hours, 0) DESC, wo.downstart DESC, wo.wo DESC
        """,
        (range_start, range_end, *filter_params),
    )
    return [
        {
            "wo": _text(row.get("wo")),
            "date_wo": _day(row.get("date_wo")),
            "hours": round(_num(row.get("hours")), 2) if row.get("hours") is not None else None,
            "part_cost": round(_num(row.get("part_cost")), 2),
            "downstart": _day(row.get("downstart")),
            "downend": _day(row.get("downend")),
            "region": _text(row.get("region_name")) or "Unassigned",
            "tribe": _text(row.get("tribe_name")) or "Unassigned",
            "state": _text(row.get("state_name")) or "Unassigned",
            "casino": _text(row.get("casino_name")) or "Unassigned",
            "serial": _text(row.get("serial")),
            "vendor": _text(row.get("vendor_name")),
            "cabinet": _text(row.get("cabinet_name")),
            "theme": _text(row.get("theme_name")),
        }
        for row in rows
    ]


def _export_workbook_bytes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Work orders"
    ws.append(list(_EXPORT_HEADERS))
    for row in rows:
        ws.append(
            [
                row.get("wo"),
                row.get("date_wo"),
                row.get("hours"),
                row.get("part_cost"),
                row.get("downstart"),
                row.get("downend"),
                row.get("region"),
                row.get("tribe"),
                row.get("state"),
                row.get("casino"),
                row.get("serial"),
                row.get("vendor"),
                row.get("cabinet"),
                row.get("theme"),
            ]
        )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.get("/export")
def capacity_export(
    start: Annotated[str | None, Query()] = None,
    end: Annotated[str | None, Query()] = None,
    month: Annotated[str | None, Query()] = None,
    region: Annotated[str | None, Query()] = None,
    tribe: Annotated[str | None, Query()] = None,
    state: Annotated[str | None, Query()] = None,
    casino: Annotated[str | None, Query()] = None,
    vendor: Annotated[str | None, Query()] = None,
    cabinet: Annotated[str | None, Query()] = None,
    theme: Annotated[str | None, Query()] = None,
    user: Annotated[dict[str, Any] | None, Depends(require_demo_user)] = None,
):
    """Download work-order line data for the current filters. Date span may exceed the dashboard 365-day cap."""
    del user
    range_start, range_end, _month_key = _resolve_range(
        start, end, month, max_span_days=_MAX_EXPORT_SPAN_DAYS
    )
    _active_filters, filter_clause, filter_params = _query_filter_params(
        region=region,
        tribe=tribe,
        state=state,
        casino=casino,
        vendor=vendor,
        cabinet=cabinet,
        theme=theme,
    )
    try:
        rows = _line_rows_for_export(
            range_start,
            range_end,
            filter_clause=filter_clause,
            filter_params=filter_params,
        )
        content = _export_workbook_bytes(rows)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"database error: {exc}") from exc

    filename = f"capacity-work-orders-{range_start.isoformat()}-to-{range_end.isoformat()}.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
